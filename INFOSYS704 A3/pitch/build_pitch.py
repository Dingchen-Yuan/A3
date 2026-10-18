#!/usr/bin/env python3
"""Five-minute Managing Director pitch. Disclaimer slide is last."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree
import copy

RED = RGBColor(0xE1, 0x06, 0x00)
INK = RGBColor(0x1C, 0x1B, 0x1A)
MUTED = RGBColor(0x5E, 0x5A, 0x55)
PAPER = RGBColor(0xF6, 0xF3, 0xEF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OK = RGBColor(0x0D, 0x6A, 0x45)
LINE = RGBColor(0xE6, 0xE0, 0xD8)

W, H = Inches(13.333), Inches(7.5)
PAGES = 12  # numbered slides; the disclaimer is not numbered
IMG = __import__("os").path.join(__import__("os").path.dirname(__import__("os").path.abspath(__file__)), "img")


def set_run(run, text, size, color, bold=False):
    run.text = text
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold


def add_text(slide, l, t, w, h, text, size, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    set_run(p.add_run(), text, size, color, bold)
    return box


def add_para(tf, text, size, color=INK, bold=False, space_before=0, space_after=0):
    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    set_run(p.add_run(), text, size, color, bold)
    return p


def rect(slide, l, t, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def chrome(slide, page, pages=PAGES):
    rect(slide, 0, 0, W, Inches(0.08), RED)
    add_text(slide, Inches(0.55), Inches(7.12), Inches(8), Inches(0.28),
             "Kmart Delivery Assist  ·  Managing Director", 12, MUTED)
    add_text(slide, Inches(10.4), Inches(7.12), Inches(2.4), Inches(0.28),
             f"{page}  /  {pages}", 12, MUTED, align=PP_ALIGN.RIGHT)


def picture(slide, name, l, t, w=None, h=None):
    import os
    return slide.shapes.add_picture(os.path.join(IMG, name), l, t, w, h)


def frame(slide, l, t, w, h):
    """Thin border behind a screenshot."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    shape.fill.background()
    shape.line.color.rgb = LINE
    shape.line.width = Pt(1)
    return shape


def new(prs, bg=WHITE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, W, H, bg)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "Kmart Delivery Assist"
    prs.core_properties.subject = "INFOSYS 704 Lab 12 prototype pitch"

    # 1 Title
    s = new(prs, INK)
    rect(s, 0, 0, Inches(0.18), H, RED)
    add_text(s, Inches(0.85), Inches(1.55), Inches(10), Inches(0.4), "KMART", 18, RED, True)
    add_text(s, Inches(0.85), Inches(2.05), Inches(11), Inches(1.4), "Delivery Assist", 60, WHITE, True)
    add_text(s, Inches(0.85), Inches(3.55), Inches(10), Inches(1.1),
             "See the order. Understand the problem.\nChoose the recovery. Finish it.", 24, RGBColor(0xE6, 0xE0, 0xD8))
    add_text(s, Inches(0.85), Inches(5.7), Inches(10), Inches(0.8),
             "Dingchen Yuan   ·   Taariq Mohamed   ·   Huy Nguyen\nWES  ·  SAP Build 13  ·  five minutes", 16, RGBColor(0xC8, 0xC2, 0xBA))
    notes(s, "BARRY · 0:00 to 0:15. Good morning. We are here to show the customer face of delivery recovery. The app is open. I will spend most of the five minutes inside it. Title only. Do not read the subtitle.")

    # 2 Situation
    s = new(prs)
    chrome(s, 2)
    add_text(s, Inches(0.55), Inches(0.35), Inches(4), Inches(0.3), "SITUATION", 14, RED, True)
    add_text(s, Inches(0.55), Inches(0.75), Inches(12), Inches(1.3),
             "The order is already in the business.\nThe customer still cannot see one story.", 32, INK, True)
    points = [
        ("Capital is in delivery", "CFCs and the online order system are there to make delivery faster and cheaper."),
        ("The order has a record", "A sales order, an outbound delivery and a freight order already describe the journey."),
        ("The customer gets fragments", "A carrier page, then a phone call, then a refund somewhere else."),
    ]
    for i, (title, body) in enumerate(points):
        left = Inches(0.55 + i * 4.15)
        rect(s, left, Inches(3.15), Inches(3.9), Inches(3.3), PAPER)
        add_text(s, left + Inches(0.25), Inches(3.4), Inches(3.4), Inches(1.3), title, 22, INK, True)
        add_text(s, left + Inches(0.25), Inches(4.8), Inches(3.4), Inches(1.3), body, 16, MUTED)
    notes(s, "BARRY · 0:15 to 0:40. Situation only. Kmart has already invested in fulfilment. The sales order, the outbound delivery and the freight order exist. What the customer meets is still a carrier page and a call. Stop. Do not mention the matrix yet.")

    # 3 Complication
    s = new(prs)
    chrome(s, 3)
    add_text(s, Inches(0.55), Inches(0.35), Inches(4), Inches(0.3), "COMPLICATION", 14, RED, True)
    add_text(s, Inches(0.55), Inches(0.75), Inches(12), Inches(1.5),
             "A failed delivery becomes an exception\nnobody closes with the customer.", 32, INK, True)
    rows = [
        ("The customer", "Cannot tell a delay from a failed attempt, or choose the next step."),
        ("The order", "Waits. A redelivery, a collection, a replacement or a refund are separate queues."),
        ("The result", "Service slips, and a completed order — the outcome the programme protects — is lost."),
    ]
    for i, (who, what) in enumerate(rows):
        top = Inches(2.7 + i * 1.25)
        add_text(s, Inches(0.55), top, Inches(3.2), Inches(0.8), who, 22, RED, True)
        add_text(s, Inches(3.8), top, Inches(8.6), Inches(0.9), what, 22, INK)
    notes(s, "BARRY · 0:40 to 1:05. Complication. Say the three lines slowly. End on this sentence: the customer cannot finish the recovery, so the order stays an exception. Then go to the resolution.")

    # 4 Resolution
    s = new(prs)
    chrome(s, 4)
    add_text(s, Inches(0.55), Inches(0.35), Inches(4), Inches(0.3), "RESOLUTION", 14, RED, True)
    add_text(s, Inches(0.55), Inches(0.7), Inches(12), Inches(1.0),
             "One Kmart journey. One service request.", 32, INK, True)
    steps = [
        "Track my\norder",
        "Understand\nwhat is happening",
        "Identify the\nproblem",
        "Choose\nwhat to do",
        "Track the\nrecovery",
        "Complete\nrefund/replacement",
    ]
    for i, label in enumerate(steps):
        left = Inches(0.5 + i * 2.12)
        rect(s, left, Inches(2.5), Inches(1.95), Inches(2.5), RED if i == 0 else PAPER)
        add_text(s, left + Inches(0.12), Inches(2.65), Inches(1.7), Inches(0.4), f"0{i+1}", 14, WHITE if i == 0 else RED, True)
        add_text(s, left + Inches(0.12), Inches(3.15), Inches(1.7), Inches(1.5), label, 16, WHITE if i == 0 else INK, True)
    add_text(s, Inches(0.55), Inches(5.4), Inches(12), Inches(1.2),
             "The customer stays in Kmart. The case is a SAP service request on the sales order.\nDemo order KM185002 — the damaged storage box.", 20, MUTED)
    notes(s, "BARRY · 1:05 to 1:30. Hand over to Huy at the end. Read the six steps only as a path, not as a list to memorise. Point at the screen if the app is already projected. Tell them the demo order is KM185002, the damaged storage box. The cushion in the same order is fine, so a refund is eighteen dollars, not thirty.")

    # 5 Why this solution
    s = new(prs)
    chrome(s, 5)
    add_text(s, Inches(0.55), Inches(0.35), Inches(8), Inches(0.3), "WHY THIS SOLUTION", 14, RED, True)
    add_text(s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.9),
             "A2 chose the data foundation first.\nThis pool can run the customer recovery.", 28, INK, True)
    cells = [
        (0, 0, PAPER, "A2, do first", "P2 decision data\nWeighted 4.50.\nNeeds OMS, WMS, TMS and Finance."),
        (1, 0, RED, "Prototype now", "Delivery Assist\nSame exception goal.\nSales order, delivery, freight, service request."),
        (0, 1, WHITE, "A2, later", "P1 routing automation\nWeighted 4.10.\nOnly after shadow evidence."),
        (1, 1, PAPER, "Not enough", "A status page with no recovery\nLow effort, low impact."),
    ]
    for col, row, fill, kicker, body in cells:
        left = Inches(0.55 + col * 6.2)
        top = Inches(1.8 + row * 2.35)
        rect(s, left, top, Inches(5.95), Inches(2.2), fill)
        title_color = WHITE if fill == RED else RED
        body_color = WHITE if fill == RED else INK
        add_text(s, left + Inches(0.3), top + Inches(0.2), Inches(5.3), Inches(0.35), kicker.upper(), 13, title_color, True)
        add_text(s, left + Inches(0.3), top + Inches(0.6), Inches(5.3), Inches(1.4), body, 20, body_color, True)
    notes(s, "HUY · 1:30 to 1:55. Start from the A2 effort-impact result, do not invent a new winner. P2 scored 4.50 and stays the programme. P1 scored 4.10 and waits for shadow evidence. Say why the prototype is Delivery Assist: the API pool has no routing service and no cost dictionary, so the internal foundation cannot be built here. The pool can read the order journey and open a service request. That is the exception the customer feels. A status page with no choice does not recover it.")

    # 6 Design and the API change
    s = new(prs)
    chrome(s, 6)
    add_text(s, Inches(0.55), Inches(0.35), Inches(8), Inches(0.3), "DESIGN  ·  OPTION 3", 14, RED, True)
    add_text(s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.6),
             "The solution changed. The new models are in the prototype.", 26, INK, True)
    stages = [
        ("Empathise", "A failed delivery is a dead end."),
        ("Define", "One job: see it, choose, finish."),
        ("Ideate", "Picked the recovery journey."),
        ("Prototype", "Six steps, one case."),
        ("Test", "Refund asks twice. A healthy order stays healthy."),
    ]
    for i, (name, line) in enumerate(stages):
        top = Inches(1.55 + i * 0.72)
        add_text(s, Inches(0.55), top, Inches(2.3), Inches(0.5), name, 18, RED, True)
        add_text(s, Inches(2.9), top, Inches(4.3), Inches(0.5), line, 18, INK)
    rect(s, Inches(7.5), Inches(1.6), Inches(5.3), Inches(4.7), PAPER)
    add_text(s, Inches(7.75), Inches(1.8), Inches(4.9), Inches(0.4), "NEW MODELS", 13, RED, True)
    add_text(s, Inches(7.75), Inches(2.25), Inches(4.9), Inches(3.8),
             "AS-IS\nCarrier page, then a phone call, then a separate refund queue.\n\nTO-BE\nOne Kmart journey. One service request.\n\nNot available\nThe billing API cannot create a refund. The app confirms it and links the sales order.",
             16, INK)
    notes(s, "HUY · 1:55 to 2:20. This is option 3, not a small wrapper change. Name the five stages once: empathise, define, ideate, prototype, test. Then the models. AS-IS file is design/A3_AS-IS_Customer_Exception.bpmn. TO-BE file is design/A3_TO-BE_Delivery_Assist.bpmn. Say the billing API can read, reject or delete, and cannot create a refund. Do not claim a live refund posting. The write we do have is the service request.")

    # 7 Usability heuristics
    s = new(prs)
    chrome(s, 7)
    add_text(s, Inches(0.55), Inches(0.35), Inches(8), Inches(0.3), "EASE OF USE", 14, RED, True)
    add_text(s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.55),
             "Nielsen’s ten heuristics, each tied to a screen.", 26, INK, True)
    heuristics = [
        ("1  Status", "Step x of 6 stays on screen."),
        ("2  Real world", "“Track my order”, not SAP codes."),
        ("3  Control", "Back, and start the order again."),
        ("4  Consistency", "One button, one step pattern."),
        ("5  Prevention", "A refund is named, then confirmed."),
        ("6  Recognition", "Recent orders are visible."),
        ("7  Flexibility", "Order number or sales order."),
        ("8  Aesthetic", "One decision on each screen."),
        ("9  Recovery", "A missing order says what to check."),
        ("10  Help", "Each screen says what happens next."),
    ]
    for i, (title, body) in enumerate(heuristics):
        col = i % 2
        row = i // 2
        left = Inches(0.55 + col * 6.3)
        top = Inches(1.55 + row * 1.0)
        add_text(s, left, top, Inches(2.4), Inches(0.7), title, 16, RED, True)
        add_text(s, left + Inches(2.4), top, Inches(3.6), Inches(0.7), body, 16, INK)
    notes(s, "HUY · 2:20 to 2:35. Do not read all ten. Say: status stays visible, the words are the customer’s, a refund asks twice, and a healthy order is not forced into a problem. The other six are on the slide if asked. Hand over to Taariq.")

    # 8 Under the hood: SAP
    s = new(prs)
    chrome(s, 8)
    add_text(s, Inches(0.55), Inches(0.35), Inches(8), Inches(0.3), "UNDER THE HOOD  ·  SAP", 14, RED, True)
    add_text(s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.6),
             "Real SAP data in. Exact SAP requests out.", 28, INK, True)
    cols = [
        ("READ LIVE", RED,
         "Sales Order API\nAPI_SALES_ORDER_SRV\nOrder, items and delivery status ($top, $select, $expand).\n\nPurchase Order API (OData V4)\nSupplier orders, items, schedule-line dates."),
        ("WRITE  ·  SIMULATED", INK,
         "Service request\nOP_API_SERVICE_REQUEST_SRV_0001\n\nPurchase-order items, schedule lines, notes.\n\nThe sandbox is read-only (POST returns 405), so each request is built in exact OData format, saved and shown on screen. One setting switches to live S/4HANA."),
        ("BUILT THE LAB WAY", INK,
         "SAP Integration Suite: API proxy and MCP server deployed.\n\nOpenCode prompts with the API spec, sandbox URL and API key.\n\nIf the sandbox is down, the app falls back to the last saved SAP snapshot."),
    ]
    for i, (kicker, color, body) in enumerate(cols):
        left = Inches(0.55 + i * 4.15)
        rect(s, left, Inches(1.7), Inches(3.9), Inches(4.95), PAPER)
        rect(s, left, Inches(1.7), Inches(3.9), Inches(0.08), color)
        add_text(s, left + Inches(0.25), Inches(1.95), Inches(3.4), Inches(0.4), kicker, 14, color, True)
        add_text(s, left + Inches(0.25), Inches(2.45), Inches(3.5), Inches(4.0), body, 14, INK)
    notes(s, "TAARIQ · 2:35 to 2:55. Twenty seconds. Say: the order data is read live from the SAP sandbox with our API key. The sandbox does not allow writes, so every write is built in the exact SAP format and shown in the app, ready for production. We deployed the API proxy and MCP server in Integration Suite and built the app with OpenCode prompts. Then go to the app.")

    # 9 Demo
    s = new(prs, INK)
    rect(s, 0, 0, Inches(0.18), H, RED)
    add_text(s, Inches(0.85), Inches(0.45), Inches(6), Inches(0.3), "SHOW IT", 14, RED, True)
    add_text(s, Inches(0.85), Inches(1.0), Inches(11), Inches(0.8), "Order KM185002", 48, WHITE, True)
    add_text(s, Inches(0.85), Inches(2.0), Inches(10), Inches(0.5),
             "ANKO Storage Box 30L  ·  damaged on delivery  ·  cushion is fine", 20, RGBColor(0xE6, 0xE0, 0xD8))
    script = [
        "Open the order. Read the one-sentence status.",
        "Confirm the problem. Keep the recommended replacement.",
        "Check for an update until the case is ready.",
        "Confirm. Show the SAP service request.",
    ]
    for i, line in enumerate(script):
        add_text(s, Inches(0.85), Inches(3.1 + i * 0.7), Inches(11), Inches(0.55), f"{i+1}    {line}", 22, WHITE)
    notes(s, "TAARIQ · 2:55 to 4:05. Leave this slide only if the app is not yet on screen. Preferred: the app is already full screen and you do not linger here. Clicks: Storage box. See the problem. This is the problem. Continue, replacement is already selected. Check for an update. Check again. Review the outcome. Confirm replacement. Open the SAP service request and point at the sales order 450218 and the service request API. If the demo fails, skip to the backup slide (12) and talk through the screenshots. Do not hunt for menus.")

    # 10 Supply Hub: the buyer side of the same exceptions
    s = new(prs)
    chrome(s, 10)
    add_text(s, Inches(0.55), Inches(0.35), Inches(9), Inches(0.3), "SAME EXCEPTIONS  ·  BUYER SIDE", 14, RED, True)
    add_text(s, Inches(0.55), Inches(0.7), Inches(12.3), Inches(0.6),
             "Before an order reaches a customer, the stock has to arrive.", 28, INK, True)
    frame(s, Inches(0.55), Inches(1.6), Inches(7.1), Inches(4.2))
    picture(s, "supplyhub_tracker.png", Inches(0.56), Inches(1.61), Inches(7.08))
    add_text(s, Inches(0.55), Inches(5.9), Inches(7.1), Inches(0.4),
             "Supply Hub · Delivery Tracker · Knit Throw, 9 days late (live SAP purchase order)", 12, MUTED)
    hub = [
        ("Delivery Tracker", "Inbound supplier orders on a map. SAP schedule lines flag a delay, short shipment, damage or supplier rejection."),
        ("Same recovery path", "Expedite, replace, refund or re-source, each built as SAP purchase-order changes."),
        ("Mix-Lock", "Defers the Anko colour and size split until early sell-through is in. Knit Throw markdown risk: $2,218 to $605."),
    ]
    for i, (title, body) in enumerate(hub):
        top = Inches(1.6 + i * 1.45)
        add_text(s, Inches(8.0), top, Inches(4.8), Inches(0.4), title, 18, RED, True)
        add_text(s, Inches(8.0), top + Inches(0.42), Inches(4.8), Inches(0.95), body, 15, INK)
    notes(s, "TAARIQ · 4:05 to 4:30. Twenty-five seconds, no demo unless asked. Say: Delivery Assist is the customer face. Supply Hub, in the same app, is the buyer side. It catches the same kinds of delivery problems earlier, on the supplier purchase order, using the same SAP API. Mix-Lock reduces the markdown risk before stock is even shipped. Hand back to Barry for the ask.")

    # 11 Close
    s = new(prs)
    chrome(s, 11)
    add_text(s, Inches(0.55), Inches(0.35), Inches(4), Inches(0.3), "THE ASK", 14, RED, True)
    add_text(s, Inches(0.55), Inches(0.8), Inches(12), Inches(1.2),
             "Use this as the customer face\nof exception recovery.", 34, INK, True)
    asks = [
        ("Today", "A customer can finish a damaged, missed or missing delivery without a call."),
        ("Next", "Supply Hub brings the same recovery to inbound supplier orders. Decision data and routing stay on the programme."),
        ("Guardrail", "A refund is explicit. A healthy order is not turned into a case."),
    ]
    for i, (title, body) in enumerate(asks):
        top = Inches(2.7 + i * 1.25)
        add_text(s, Inches(0.55), top, Inches(2.4), Inches(0.7), title, 22, RED, True)
        add_text(s, Inches(3.1), top, Inches(9.3), Inches(0.9), body, 22, INK)
    notes(s, "BARRY · 4:30 to 4:50. Close on the ask, not a summary of every slide. Today the customer can finish the recovery. The heavier routing work stays later. A refund is a decision, not a side effect. Thank them and take one question. Then the disclaimer slide is for the upload, not for a speech.")

    # 12 Backup screenshots, in case the live demo fails
    s = new(prs)
    chrome(s, 12)
    add_text(s, Inches(0.55), Inches(0.35), Inches(9), Inches(0.3), "IF THE DEMO FAILS  ·  ORDER KM185002", 14, RED, True)
    add_text(s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.6),
             "The same journey, screen by screen.", 26, INK, True)
    shots = [
        ("assist_1_understand.png", "2  Understand"),
        ("assist_2_problem.png", "3  Identify the problem"),
        ("assist_3_choose.png", "4  Choose what to do"),
        ("assist_4_track.png", "5  Track the recovery"),
        ("assist_5_complete.png", "6  Complete, with the SAP request"),
    ]
    for i, (name, label) in enumerate(shots):
        left = Inches(0.55 + i * 2.5)
        frame(s, left, Inches(1.55), Inches(2.3), Inches(5.0))
        pic = picture(s, name, left + Inches(0.01), Inches(1.56), Inches(2.28))
        # keep the phone inside the frame: crop the bottom if it runs long
        max_h = Inches(4.98)
        if pic.height > max_h:
            pic.crop_bottom = 1 - max_h / pic.height
            pic.height = max_h
        add_text(s, left, Inches(6.62), Inches(2.3), Inches(0.4), label, 12, MUTED)
    notes(s, "Backup only. Use it if the app will not load in the lab. Talk through the five screens in under a minute, then go to the ask.")

    # 13 Disclaimer — must be last
    s = new(prs)
    rect(s, 0, 0, W, Inches(0.08), RED)
    add_text(s, Inches(0.4), Inches(0.2), Inches(12.5), Inches(0.4),
             "GROUPWORK MANAGEMENT DISCLAIMER", 22, INK, True)
    add_text(s, Inches(0.4), Inches(0.65), Inches(12.5), Inches(0.85),
             "We take equal contribution in groupwork seriously, and marks may be deducted where there is a significant inequality in the team’s contribution. In the following table, state each group member’s contribution* to this assignment. You may also include any comments to support the % contribution indicated on the table.",
             13, INK)
    add_text(s, Inches(0.4), Inches(1.5), Inches(12.5), Inches(0.35),
             "*No contribution to the deliverable may result in a deduction of marks in the group assignment.",
             12, MUTED)

    rows_h = [Inches(0.32), Inches(0.38), Inches(0.38), Inches(0.55)]
    table = s.shapes.add_table(4, 5, Inches(0.4), Inches(1.95), Inches(12.5), Inches(1.7)).table
    table.columns[0].width = Inches(2.3)
    for i in range(1, 5):
        table.columns[i].width = Inches(2.55)
    data = [
        ["GROUP MEMBER", "Dingchen Yuan", "Taariq Mohamed", "Huy Nguyen", "Total"],
        ["UPI", "dyua155", "tmoh716", "hugn456", ""],
        ["% Contribution", "34%", "33%", "33%", "100%"],
        ["Comments", "Equal split", "Equal split", "Equal split", "Rounding"],
    ]
    for r, row in enumerate(data):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER if c else PP_ALIGN.LEFT
                for run in p.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(12)
                    run.font.bold = r == 0 or c == 0
                    run.font.color.rgb = WHITE if r == 0 else INK
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            fill = cell.fill
            fill.solid()
            fill.fore_color.rgb = RED if r == 0 else (PAPER if r % 2 == 0 else WHITE)

    add_text(s, Inches(0.4), Inches(3.85), Inches(12.5), Inches(1.15),
             "Criteria explanation:\n1. UPI: Enter each team member’s UPI.\n2. % Contribution: Enter a percentage for each member. The figures must add up to 100%.\n3. Comments (optional): Add a comment where the percentage needs context.",
             13, INK)
    notes(s, "Do not spend speaking time here unless asked. This slide has to be last in the file. UPIs are filled in. Confirm the 34/33/33 split with the whole group. If the split changes, change the percentages so they still total 100.")

    out = "pitch/Kmart-Delivery-Assist-Lab12.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    build()
